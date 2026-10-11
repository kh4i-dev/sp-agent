#!/usr/bin/env python3
"""Tiêm ghi chú diễn giả (Speaker Notes) từ file Markdown vào file PowerPoint (.pptx).

Định dạng file Markdown (speaker_notes.md):
    <!-- slide 1 | Tiêu đề slide 1 -->
    Nội dung lời thoại slide 1...

    <!-- slide 2 | Tiêu đề slide 2 -->
    Nội dung lời thoại slide 2...

Cách dùng:
    python add_notes.py deck.pptx notes.md               # Tạo bản copy deck_notes.pptx
    python add_notes.py deck.pptx notes.md --in-place    # Ghi đè trực tiếp lên deck.pptx
    python add_notes.py deck.pptx notes.md -o output.pptx # Đặt tên file xuất cụ thể
"""
import argparse
import re
import sys
from pathlib import Path

# Đảm bảo in tiếng Việt không bị lỗi charmap trên Windows terminal
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def load_notes_from_md(md_path: Path) -> dict[int, str]:
    """Phân tích cú pháp file markdown chứa speaker notes."""
    text = md_path.read_text(encoding="utf-8")
    # Tách theo mẫu <!-- slide 1 | ... --> hoặc <!-- slide 1 -->
    pattern = r"<!--\s*slide\s+(\d+)(?:\s*\|.*?)?\s*-->"
    parts = re.split(pattern, text, flags=re.IGNORECASE)
    notes: dict[int, str] = {}
    for i in range(1, len(parts) - 1, 2):
        slide_no = int(parts[i])
        content = parts[i + 1].strip()
        notes[slide_no] = content
    return notes


def inject_notes_python_pptx(pptx_path: Path, notes_map: dict[int, str], out_path: Path):
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    total_slides = len(prs.slides)

    updated_count = 0
    for idx, slide in enumerate(prs.slides, start=1):
        if idx in notes_map and notes_map[idx]:
            notes_slide = slide.notes_slide
            tf = notes_slide.notes_text_frame
            tf.text = notes_map[idx]
            updated_count += 1

    prs.save(str(out_path))
    return updated_count, total_slides


def main():
    ap = argparse.ArgumentParser(description="Tiêm speaker notes từ markdown vào PPTX.")
    ap.add_argument("pptx", help="Đường dẫn file .pptx nguồn")
    ap.add_argument("notes_md", help="Đường dẫn file markdown chứa speaker notes")
    ap.add_argument("-o", "--output", help="Đường dẫn file .pptx xuất (mặc định: <name>_notes.pptx)")
    ap.add_argument("--in-place", action="store_true", help="Ghi đè trực tiếp lên file pptx gốc")
    args = ap.parse_args()

    pptx_path = Path(args.pptx)
    notes_path = Path(args.notes_md)

    if not pptx_path.exists():
        sys.exit(f"Lỗi: Không tìm thấy file PPTX: {pptx_path}")
    if not notes_path.exists():
        sys.exit(f"Lỗi: Không tìm thấy file Markdown: {notes_path}")

    if args.in-place:
        out_path = pptx_path
    elif args.output:
        out_path = Path(args.output)
    else:
        out_path = pptx_path.with_name(f"{pptx_path.stem}_notes.pptx")

    notes_map = load_notes_from_md(notes_path)
    print(f">> Đã đọc {len(notes_map)} ghi chú từ {notes_path.name}")

    try:
        updated, total = inject_notes_python_pptx(pptx_path, notes_map, out_path)
        print(f">> Thành công: Đã tiêm {updated}/{total} slide vào: {out_path}")
    except ImportError:
        sys.exit(
            "Lỗi: Cần thư viện python-pptx để tiêm notes.\n"
            "Cài đặt: pip install python-pptx hoặc chạy qua: uv run --with python-pptx python add_notes.py ..."
        )
    except Exception as e:
        sys.exit(f"Lỗi khi tiêm notes: {e}")


if __name__ == "__main__":
    main()
